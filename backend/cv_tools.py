import json
import re
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from ai_client import generate_text


PROFILE_SCHEMA = {
    "name": "",
    "headline": "",
    "email": "",
    "phone": "",
    "location": "",
    "website": "",
    "linkedin": "",
    "summary": "",
    "skills": [],
    "experience": [
        {
            "company": "",
            "role": "",
            "location": "",
            "start_date": "",
            "end_date": "",
            "current": False,
            "bullets": [],
        }
    ],
    "education": [
        {
            "institution": "",
            "degree": "",
            "field": "",
            "start_date": "",
            "end_date": "",
        }
    ],
    "projects": [
        {"name": "", "description": "", "link": "", "technologies": []}
    ],
    "certifications": [],
    "languages": [],
    "custom_sections": {},
}


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return ""
    return str(value).strip()


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [_clean(item) for item in value if _clean(item)]


def normalize_profile(raw: Any) -> dict[str, Any]:
    data = raw if isinstance(raw, dict) else {}
    result: dict[str, Any] = {
        "name": _clean(data.get("name")),
        "headline": _clean(data.get("headline")),
        "email": _clean(data.get("email")),
        "phone": _clean(data.get("phone")),
        "location": _clean(data.get("location")),
        "website": _clean(data.get("website")),
        "linkedin": _clean(data.get("linkedin")),
        "summary": _clean(data.get("summary")),
        "skills": _string_list(data.get("skills")),
        "experience": [],
        "education": [],
        "projects": [],
        "certifications": _string_list(data.get("certifications")),
        "languages": _string_list(data.get("languages")),
        "custom_sections": data.get("custom_sections") if isinstance(data.get("custom_sections"), dict) else {},
    }

    for item in data.get("experience", []) if isinstance(data.get("experience"), list) else []:
        if not isinstance(item, dict):
            continue
        result["experience"].append({
            "company": _clean(item.get("company")),
            "role": _clean(item.get("role")),
            "location": _clean(item.get("location")),
            "start_date": _clean(item.get("start_date")),
            "end_date": _clean(item.get("end_date")),
            "current": bool(item.get("current", False)),
            "bullets": _string_list(item.get("bullets")),
        })

    for item in data.get("education", []) if isinstance(data.get("education"), list) else []:
        if not isinstance(item, dict):
            continue
        result["education"].append({
            "institution": _clean(item.get("institution")),
            "degree": _clean(item.get("degree")),
            "field": _clean(item.get("field")),
            "start_date": _clean(item.get("start_date")),
            "end_date": _clean(item.get("end_date")),
        })

    for item in data.get("projects", []) if isinstance(data.get("projects"), list) else []:
        if not isinstance(item, dict):
            continue
        result["projects"].append({
            "name": _clean(item.get("name")),
            "description": _clean(item.get("description")),
            "link": _clean(item.get("link")),
            "technologies": _string_list(item.get("technologies")),
        })

    return result


def _extract_json(raw: str) -> dict[str, Any]:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
        text = re.sub(r"\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            return json.loads(text[start : end + 1])
        raise ValueError("AI returned invalid profile JSON.")


def build_profile_from_text(source_text: str, source_type: str = "cv") -> dict[str, Any]:
    if not source_text or len(source_text.strip()) < 20:
        raise ValueError("The source text is too short to build a profile.")

    prompt = f"""
You are a careful CV/profile extraction engine.
Create a structured professional profile from the source below.

SOURCE TYPE: {source_type}

STRICT RULES:
- Use ONLY facts explicitly supported by the source.
- Never invent employers, dates, degrees, skills, achievements, contact details, URLs, or metrics.
- Preserve names and dates as written when possible.
- If information is missing, use an empty string or empty array.
- Convert work/project descriptions into concise bullet points only when the source supports them.
- Do not add generic filler achievements.
- Return ONLY valid JSON matching the schema.

SCHEMA:
{json.dumps(PROFILE_SCHEMA, ensure_ascii=False, indent=2)}

SOURCE:
{source_text[:50000]}
"""

    raw = generate_text(prompt, json_output=True)
    return normalize_profile(_extract_json(raw))


def analyze_ats(profile: dict[str, Any], job_description: str = "") -> dict[str, Any]:
    clean_profile = normalize_profile(profile)
    prompt = f"""
You are an ATS and CV quality reviewer.
Review the candidate profile below.

If a job description is supplied, evaluate alignment against it. If it is empty,
perform a general ATS-readiness review.

Do not invent facts. Do not penalize the candidate for information that simply
isn't present unless it is relevant to ATS completeness. Keep recommendations
specific and actionable.

Return ONLY JSON with this shape:
{{
  "score": 0,
  "summary": "",
  "strengths": [],
  "issues": [],
  "keywords_found": [],
  "keywords_missing": [],
  "recommendations": []
}}

PROFILE:
{json.dumps(clean_profile, ensure_ascii=False, indent=2)}

JOB DESCRIPTION:
{job_description[:30000]}
"""
    raw = generate_text(prompt, json_output=True)
    data = _extract_json(raw)
    try:
        score = max(0, min(100, int(float(data.get("score", 0)))))
    except (TypeError, ValueError):
        score = 0
    return {
        "score": score,
        "summary": _clean(data.get("summary")),
        "strengths": _string_list(data.get("strengths")),
        "issues": _string_list(data.get("issues")),
        "keywords_found": _string_list(data.get("keywords_found")),
        "keywords_missing": _string_list(data.get("keywords_missing")),
        "recommendations": _string_list(data.get("recommendations")),
    }


def create_cv_docx(profile: dict[str, Any]) -> bytes:
    """Create a DOCX CV using the selected visual template."""
    p = normalize_profile(profile)
    template = str(profile.get("template", "ats-classic") or "ats-classic").strip().lower()
    if template not in {"ats-classic", "modern", "executive", "minimal"}:
        template = "ats-classic"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5 if template != "minimal" else 9)

    # Template-specific visual system. All templates remain text-based and
    # editable in Word; no screenshots or flattened CV pages are generated.
    from docx.shared import RGBColor

    configs = {
        "ats-classic": {"accent": RGBColor(15, 23, 42), "name_size": 20, "heading_size": 10.5, "center": True, "rule": False},
        "modern": {"accent": RGBColor(2, 132, 199), "name_size": 22, "heading_size": 10.5, "center": False, "rule": True},
        "executive": {"accent": RGBColor(30, 41, 59), "name_size": 21, "heading_size": 11, "center": True, "rule": True},
        "minimal": {"accent": RGBColor(71, 85, 105), "name_size": 19, "heading_size": 9.5, "center": False, "rule": False},
    }
    cfg = configs[template]

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER if cfg["center"] else WD_ALIGN_PARAGRAPH.LEFT
    name.paragraph_format.space_after = Pt(1)
    run = name.add_run(p["name"] or "Your Name")
    run.bold = True
    run.font.size = Pt(cfg["name_size"])
    run.font.color.rgb = cfg["accent"]

    if p["headline"]:
        head = doc.add_paragraph()
        head.alignment = name.alignment
        head.paragraph_format.space_after = Pt(2)
        run = head.add_run(p["headline"])
        run.bold = True
        run.font.size = Pt(10.5)

    contact = [p["email"], p["phone"], p["location"], p["website"], p["linkedin"]]
    contact = [x for x in contact if x]
    if contact:
        para = doc.add_paragraph()
        para.alignment = name.alignment
        para.paragraph_format.space_after = Pt(5)
        para.add_run("  |  ".join(contact)).font.size = Pt(8.5)


    def heading(title: str):
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(7 if template != "minimal" else 5)
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(cfg["heading_size"])
        run.font.color.rgb = cfg["accent"]
        return para

    if p["summary"]:
        heading("Professional Summary")
        doc.add_paragraph(p["summary"])

    if p["skills"]:
        heading("Skills")
        doc.add_paragraph(" • ".join(p["skills"]))

    if p["experience"]:
        heading("Experience")
        for item in p["experience"]:
            if not any(item.values()):
                continue
            title = doc.add_paragraph()
            title.paragraph_format.space_after = Pt(0)
            role = item["role"] or "Role"
            company = item["company"]
            r = title.add_run(f"{role}{' — ' + company if company else ''}")
            r.bold = True
            
            if template == "executive":
                r.font.color.rgb = cfg["accent"]
            dates = " — ".join(x for x in [item["start_date"], item["end_date"] or ("Present" if item["current"] else "")] if x)
            if dates:
                title.add_run(f"  |  {dates}")
            if item["location"]:
                title.add_run(f"  |  {item['location']}")
            for bullet in item["bullets"]:
                doc.add_paragraph(bullet, style="List Bullet")

    if p["education"]:
        heading("Education")
        for item in p["education"]:
            if not any(item.values()):
                continue
            title = doc.add_paragraph()
            degree = " ".join(x for x in [item["degree"], item["field"]] if x)
            title.add_run(degree or "Education").bold = True
            if item["institution"]:
                title.add_run(f" — {item['institution']}")
            dates = " — ".join(x for x in [item["start_date"], item["end_date"]] if x)
            if dates:
                title.add_run(f"  |  {dates}")

    if p["projects"]:
        heading("Projects")
        for item in p["projects"]:
            if not any(item.values()):
                continue
            title = doc.add_paragraph()
            title.add_run(item["name"] or "Project").bold = True
            if item["technologies"]:
                title.add_run(f" — {', '.join(item['technologies'])}")
            if item["description"]:
                doc.add_paragraph(item["description"])
            if item["link"]:
                doc.add_paragraph(item["link"])

    if p["certifications"]:
        heading("Certifications")
        for item in p["certifications"]:
            doc.add_paragraph(item, style="List Bullet")

    if p["languages"]:
        heading("Languages")
        doc.add_paragraph(" • ".join(p["languages"]))

    for section_name, values in p["custom_sections"].items():
        if not values:
            continue
        heading(str(section_name))
        if isinstance(values, list):
            for value in values:
                doc.add_paragraph(_clean(value), style="List Bullet")
        else:
            doc.add_paragraph(_clean(values))

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
