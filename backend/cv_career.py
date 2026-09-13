"""CV & Career helpers for MyToolbox AI.

The module keeps the CV workflow deliberately bounded:
1. extract source text from PDF/DOCX,
2. build a structured profile from source facts,
3. optionally score that profile against a job description.
"""

from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

import fitz
from docx import Document

from ai_client import generate_text
from cv_tools import analyze_ats, build_profile_from_text, normalize_profile


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return ""
    return str(value).strip()


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [_clean(item) for item in value if _clean(item)]


def _extract_json(raw: str) -> dict[str, Any]:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
        text = re.sub(r"\s*```$", "", text)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start < 0 or end <= start:
            raise ValueError("AI returned invalid CV analysis JSON.")
        data = json.loads(text[start : end + 1])
    if not isinstance(data, dict):
        raise ValueError("AI returned invalid CV analysis JSON.")
    return data


def extract_cv_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from a PDF or DOCX CV, with OCR fallback for scanned PDFs."""
    lower = filename.lower()
    if lower.endswith(".docx"):
        document = Document(BytesIO(file_bytes))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    parts.append(line)
        text = "\n".join(parts).strip()
        if len(text) < 20:
            raise ValueError("The DOCX CV contains too little readable text.")
        return text[:60000]

    if not lower.endswith(".pdf"):
        raise ValueError("Please upload a PDF or DOCX CV.")

    try:
        document = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise ValueError(f"Could not read the PDF CV: {exc}") from exc

    try:
        pages = [{"page": index + 1, "text": page.get_text("text")} for index, page in enumerate(document)]
    finally:
        document.close()

    text = "\n\n".join(str(page["text"]).strip() for page in pages if str(page["text"]).strip()).strip()
    if len(text) >= 20:
        return text[:60000]

    # Reuse the project's existing OCR implementation only when a text layer
    # is genuinely missing. This avoids making every CV extraction expensive.
    try:
        from pdf_ocr import analyze_pdf

        ocr_result = analyze_pdf(file_bytes)
        ocr_pages = ocr_result.get("pages", []) if isinstance(ocr_result, dict) else []
        ocr_text = "\n\n".join(
            str(page.get("text", "")).strip()
            for page in ocr_pages
            if isinstance(page, dict) and str(page.get("text", "")).strip()
        ).strip()
        if len(ocr_text) >= 20:
            return ocr_text[:60000]
    except Exception:
        pass

    raise ValueError("The CV contains too little readable text. Try a text-based PDF/DOCX or a clearer scan.")


def _build_action_cards(data: dict[str, Any]) -> list[dict[str, Any]]:
    recommendations = _string_list(data.get("recommendations"))
    issues = _string_list(data.get("issues"))
    cards: list[dict[str, Any]] = []
    for index, recommendation in enumerate(recommendations[:4]):
        cards.append({
            "id": f"recommendation-{index + 1}",
            "priority": "high" if index == 0 else "medium",
            "title": "Improve CV alignment",
            "issue": recommendation,
            "current_text": "",
            "suggested_text": "",
            "section": "CV",
            "rationale": recommendation,
            "can_apply": False,
        })
    if not cards:
        for index, issue in enumerate(issues[:4]):
            cards.append({
                "id": f"issue-{index + 1}",
                "priority": "medium",
                "title": "Review CV issue",
                "issue": issue,
                "current_text": "",
                "suggested_text": "",
                "section": "CV",
                "rationale": issue,
                "can_apply": False,
            })
    return cards


def analyze_cv(cv_text: str, target_job: str = "") -> dict[str, Any]:
    """Return the analysis shape expected by the existing CV workspace."""
    profile = build_profile_from_text(cv_text, "cv")
    ats = analyze_ats(profile, target_job)

    prompt = f"""
You are a conservative career analyst reviewing a candidate profile.
Use ONLY information present in the profile. Never invent achievements, years,
metrics, employers, technologies, qualifications, or skills.

Return ONLY JSON with:
{{
  "candidate_name": "",
  "headline": "",
  "target_role": "",
  "summary": "",
  "skills": [],
  "strengths": [],
  "weaknesses": [],
  "matched_keywords": [],
  "missing_keywords": [],
  "recommendations": [],
  "experience": [
    {{"role":"", "company":"", "dates":"", "bullets":[]}}
  ],
  "job_match_summary": ""
}}

PROFILE:
{json.dumps(profile, ensure_ascii=False, indent=2)}

TARGET JOB:
{target_job[:30000]}
"""
    data = _extract_json(generate_text(prompt, json_output=True))

    breakdown = {
        "role_alignment": 0,
        "skills": 0,
        "experience": 0,
        "evidence": 0,
        "keywords": 0,
    }
    if target_job.strip():
        # The ATS model gives the authoritative overall result. The component
        # values are requested separately so the UI can remain interpretable.
        breakdown_prompt = f"""
Score this candidate profile against the target job from 0 to 100 for exactly
five dimensions: role_alignment, skills, experience, evidence, keywords.
Return ONLY JSON with integer values for those five keys.
Do not invent facts.
PROFILE: {json.dumps(profile, ensure_ascii=False)}
JOB: {target_job[:30000]}
"""
        try:
            bd = _extract_json(generate_text(breakdown_prompt, json_output=True))
            for key in breakdown:
                try:
                    breakdown[key] = max(0, min(100, int(float(bd.get(key, 0)))))
                except (TypeError, ValueError):
                    pass
        except Exception:
            pass

    return {
        "candidate_name": _clean(data.get("candidate_name")) or profile.get("name", ""),
        "headline": _clean(data.get("headline")) or profile.get("headline", ""),
        "target_role": _clean(data.get("target_role")) or profile.get("headline", ""),
        "summary": _clean(data.get("summary")) or profile.get("summary", ""),
        "skills": _string_list(data.get("skills")) or profile.get("skills", []),
        "strengths": _string_list(data.get("strengths")),
        "weaknesses": _string_list(data.get("weaknesses")),
        "missing_keywords": ats["keywords_missing"],
        "matched_keywords": ats["keywords_found"],
        "recommendations": _string_list(data.get("recommendations")) or ats["recommendations"],
        "experience": data.get("experience") if isinstance(data.get("experience"), list) else [],
        "ats_score": ats["score"],
        "ats_breakdown": breakdown,
        "job_match_summary": _clean(data.get("job_match_summary")) or ats["summary"],
        "action_cards": _build_action_cards(ats),
    }


def recalculate_ats(target_job: str, analysis: dict[str, Any], applied_actions: list[Any] | None = None) -> dict[str, Any]:
    """Re-score the edited analysis/profile payload without mutating the source file."""
    if not target_job.strip():
        raise ValueError("A target job description is required for ATS recalculation.")

    # The existing frontend sends the edited analysis object. Convert the
    # available fields back into the profile schema used by the ATS reviewer.
    profile = normalize_profile({
        "name": analysis.get("candidate_name", ""),
        "headline": analysis.get("headline", ""),
        "summary": analysis.get("summary", ""),
        "skills": analysis.get("skills", []),
        "experience": analysis.get("experience", []),
    })
    ats = analyze_ats(profile, target_job)
    return {
        "ats_score": ats["score"],
        "job_match_summary": ats["summary"],
        "matched_keywords": ats["keywords_found"],
        "remaining_missing_keywords": ats["keywords_missing"],
    }
